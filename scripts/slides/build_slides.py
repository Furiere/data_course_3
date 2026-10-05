"""
Build a slide deck for each lecture from its markdown, on the course template.

    lecture-NN-*.md  ──parse──>  slides  ──render──>  slides/lecture-NN-*.pptx
                                                      (+ .pdf with --pdf)

Mapping from the markdown:
  # H1                 title slide
  ## Contents          agenda slide
  ## H2                section divider, listing its subtopics
  ### H3 / #### H4     content slides, titled with the heading
  ```sql``` + result   example slides — code and result rendered as images
  tables               native PowerPoint tables, split across slides if long
  paragraphs, lists,
  > quotes             text slides, split across slides if long

The template (pptx_template/explorers_template_1.pptx) supplies the artwork:
its slides are cloned and filled in, then the originals are removed.

Usage:  python scripts/slides/build_slides.py [--pdf] [lecture.md ...]
Needs:  pip install -r scripts/slides/requirements.txt, Google Chrome, and
        Microsoft PowerPoint for --pdf (macOS).
"""

import argparse
import copy
import hashlib
import json
import math
import re
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from pygments import highlight
from pygments.formatters import HtmlFormatter
from pygments.lexers import PostgresLexer, TextLexer

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / 'pptx_template' / 'explorers_template_1.pptx'
OUT_DIR = ROOT / 'slides'
IMG_DIR = ROOT / 'build' / 'slides'
LECTURES = [
    'lecture-01-sql-basics.md',
    'lecture-02-ddl-dml-procedures-variables.md',
]
AUTHOR = 'Denis Boldin'
COURSE = 'SQL for Data Analysts'

# Template slides, by position in the template deck.
T_TITLE, T_BLANK, T_SECTION_A, T_SECTION_B, T_TEXT, T_PICTURE, T_QA, T_THANKS = range(8)

FONT = 'Arial'
MONO = 'Consolas'
INK = RGBColor(0x00, 0x00, 0x00)
MUTED = RGBColor(0x5F, 0x6B, 0x70)
ACCENT = RGBColor(0x03, 0xCD, 0xD8)       # the template's cyan
ACCENT_TEXT = RGBColor(0x00, 0x7C, 0x85)  # the cyan, dark enough to read as text
TINT = RGBColor(0xE6, 0xF9, 0xFA)
ZEBRA = RGBColor(0xF3, 0xF6, 0xF7)

SLIDE_W, SLIDE_H = 13.333, 7.5
LEFT = 1.1                     # content left edge, as on the template
RIGHT = 12.4
BOTTOM = 6.95
BODY_PT = 18
CAPTION_PT = 16
TABLE_PT = 14


# ---------------------------------------------------------------------------
# Markdown → blocks
# ---------------------------------------------------------------------------

@dataclass
class Block:
    kind: str                  # para | list | quote | table | code | heading
    text: str = ''
    level: int = 0             # heading level
    items: list = field(default_factory=list)   # list: (indent, ordered, text)
    rows: list = field(default_factory=list)    # table: header first
    lang: str = ''
    result: list = None        # code: result table rows, header first
    start: int = 0             # list: items already shown on a previous slide

    @property
    def plain(self):
        return strip_inline(self.text)


LIST_RE = re.compile(r'^(\s*)([-*]|\d+\.)\s+(.*)$')


def split_row(line):
    """Split a markdown table row on pipes that are not inside backticks."""
    cells, cur, in_code = [], '', False
    for ch in line.strip().strip('|'):
        if ch == '`':
            in_code = not in_code
        if ch == '|' and not in_code:
            cells.append(cur.strip())
            cur = ''
        else:
            cur += ch
    cells.append(cur.strip())
    return cells


def parse_table(lines):
    rows = [split_row(l) for l in lines]
    return [r for r in rows if not all(re.fullmatch(r':?-+:?', c) for c in r)]


def starts_block(line):
    s = line.strip()
    return (not s or s.startswith(('#', '```', '|', '>', '<!--')) or s == '---'
            or LIST_RE.match(line))


def parse_blocks(md):
    lines = md.split('\n')
    blocks, i = [], 0
    while i < len(lines):
        line, s = lines[i], lines[i].strip()
        if not s or s == '---':
            i += 1
        elif s.startswith('<!--result'):
            j = i + 1
            while lines[j].strip() != '<!--/result-->':
                j += 1
            rows = parse_table([l for l in lines[i + 1:j] if l.strip().startswith('|')])
            if blocks and blocks[-1].kind == 'code':
                blocks[-1].result = rows
            i = j + 1
        elif s.startswith('<!--'):
            i += 1
        elif s.startswith('```'):
            j = i + 1
            while not lines[j].strip().startswith('```'):
                j += 1
            blocks.append(Block('code', text='\n'.join(lines[i + 1:j]).rstrip(),
                                lang=s[3:] or 'text'))
            i = j + 1
        elif s.startswith('#'):
            hashes = len(s) - len(s.lstrip('#'))
            blocks.append(Block('heading', text=s[hashes:].strip(), level=hashes))
            i += 1
        elif s.startswith('|'):
            j = i
            while j < len(lines) and lines[j].strip().startswith('|'):
                j += 1
            blocks.append(Block('table', rows=parse_table(lines[i:j])))
            i = j
        elif s.startswith('>'):
            j, paras, cur = i, [], []
            while j < len(lines) and lines[j].strip().startswith('>'):
                t = lines[j].strip()[1:].strip()
                if t:
                    cur.append(t)
                elif cur:
                    paras.append(' '.join(cur))
                    cur = []
                j += 1
            if cur:
                paras.append(' '.join(cur))
            for p in paras:
                blocks.append(Block('quote', text=p))
            i = j
        elif LIST_RE.match(line):
            items, j = [], i
            while j < len(lines):
                m = LIST_RE.match(lines[j])
                if m:
                    items.append([len(m.group(1)) // 2, m.group(2)[0].isdigit(), m.group(3).strip()])
                elif lines[j].strip() and not starts_block(lines[j]) and lines[j].startswith(' '):
                    items[-1][2] += ' ' + lines[j].strip()
                else:
                    break
                j += 1
            blocks.append(Block('list', items=items))
            i = j
        else:
            j, cur = i, []
            while j < len(lines) and not starts_block(lines[j]):
                cur.append(lines[j].strip())
                j += 1
            blocks.append(Block('para', text=' '.join(cur)))
            i = j
    return blocks


INLINE_RE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<![\w*])\*[^*\s][^*]*\*(?![\w*]))')


def parse_inline(text, bold=False, italic=False, link=None):
    """Yield (text, style) runs; style has bold/italic/code/link."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            yield text[pos:m.start()], dict(bold=bold, italic=italic, code=False, link=link)
        tok = m.group(0)
        if tok.startswith('**'):
            yield from parse_inline(tok[2:-2], True, italic, link)
        elif tok.startswith('`'):
            yield tok[1:-1], dict(bold=bold, italic=italic, code=True, link=link)
        elif tok.startswith('['):
            label, url = re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok).groups()
            yield from parse_inline(label, bold, italic, url if url.startswith('http') else None)
        else:
            yield from parse_inline(tok[1:-1], bold, True, link)
        pos = m.end()
    if pos < len(text):
        yield text[pos:], dict(bold=bold, italic=italic, code=False, link=link)


def strip_inline(text):
    return ''.join(t for t, _ in parse_inline(text))


# ---------------------------------------------------------------------------
# Lecture outline
# ---------------------------------------------------------------------------

@dataclass
class Topic:
    title: str
    kicker: str
    blocks: list


@dataclass
class Section:
    title: str
    intro: list
    topics: list


def outline(blocks):
    title = next(b.text for b in blocks if b.kind == 'heading' and b.level == 1)
    contents, sections = [], []
    sec = h3 = None
    target = None
    for b in blocks:
        if b.kind == 'heading':
            if b.level == 1:
                target = []          # the italic byline under H1 — not on slides
            elif b.level == 2:
                h3 = None
                if b.text == 'Contents':
                    sec, target = None, contents
                else:
                    sec = Section(b.text, [], [])
                    sections.append(sec)
                    target = sec.intro
            elif b.level == 3:
                h3 = b.text
                t = Topic(b.text, sec.title, [])
                sec.topics.append(t)
                target = t.blocks
            else:
                kicker = f'{sec.title}  ·  {h3}' if h3 else sec.title
                # an H3 with no text of its own is just a heading for its H4s
                if sec.topics and sec.topics[-1].title == h3 and not sec.topics[-1].blocks:
                    sec.topics.pop()
                t = Topic(b.text, kicker, [])
                sec.topics.append(t)
                target = t.blocks
        else:
            target.append(b)
    return title, contents, sections


# ---------------------------------------------------------------------------
# Example images
# ---------------------------------------------------------------------------

CSS = """
html, body { margin: 0; padding: 0; background: transparent; }
#card { display: inline-block; box-sizing: border-box; }
.code { background: #F5F7F8; display: inline-block; border: 1px solid #DDE3E6; border-radius: 10px;
        padding: 16px 24px; font: 22px/1.45 Menlo, 'SF Mono', monospace;
        white-space: pre; color: #1B1F23; }
.code b, .code strong { font-weight: 700; }
.label { font: 700 14px Arial, sans-serif; letter-spacing: .08em; color: #5F6B70;
         text-transform: uppercase; margin: 0 0 6px 2px; }
table { border-collapse: collapse; font: 19px/1.3 Arial, sans-serif; color: #111;
        background: #fff; border: 1px solid #DDE3E6; }
th { background: #03CDD8; color: #04292C; text-align: left; padding: 7px 12px;
     font-weight: 700; white-space: nowrap; }
td { padding: 6px 12px; border-top: 1px solid #E4E9EB; max-width: 440px;
     vertical-align: top; }
tr:nth-child(even) td { background: #F3F6F7; }
td.null { color: #9AA5AA; font-style: italic; }
td.num { text-align: right; font-variant-numeric: tabular-nums; }
table.plan td { font: 17px/1.35 Menlo, monospace; white-space: pre; max-width: none; }
"""


def esc(s):
    return s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def page(body):
    return f'<!doctype html><html><head><meta charset="utf-8"><style>{CSS}</style></head><body>{body}</body></html>'


def code_pane(code, lang):
    lexer = PostgresLexer() if lang == 'sql' else TextLexer()
    fmt = HtmlFormatter(nowrap=True, noclasses=True, style='friendly')
    # the lexer flags psql meta-commands (\\set) as errors; do not draw the red box
    html = highlight(code, lexer, fmt).replace('border: 1px solid #FF0000', '')
    return f'<div class="code">{html}</div>'


def code_html(code, lang):
    return page(f'<div id="card">{code_pane(code, lang)}</div>')


def code_html_split(code, lang):
    """Long code in two panes side by side, split at the blank line nearest the middle."""
    lines = code.split('\n')
    if len(lines) < 18:
        return None
    blanks = [i for i, l in enumerate(lines) if not l.strip() and 0.3 * len(lines) <= i <= 0.7 * len(lines)]
    if not blanks:
        return None
    cut = min(blanks, key=lambda i: abs(i - len(lines) / 2))
    left, right = '\n'.join(lines[:cut]), '\n'.join(lines[cut + 1:])
    return page(f'<div id="card" style="display:inline-flex;gap:20px;align-items:flex-start">'
                f'{code_pane(left, lang)}{code_pane(right, lang)}</div>')


def result_table(head, body, plan):
    th = ''.join(f'<th>{esc(h)}</th>' for h in head)
    trs = []
    for r in body:
        tds = []
        for c in r:
            cls = 'null' if c == 'NULL' else 'num' if re.fullmatch(r'-?[\d.,]+', c) else ''
            tds.append(f'<td class="{cls}">{esc(c)}</td>')
        trs.append('<tr>' + ''.join(tds) + '</tr>')
    if not body:
        trs.append(f'<tr><td class="null" colspan="{len(head)}">(no rows)</td></tr>')
    cls = ' class="plan"' if plan else ''
    return f'<table{cls}><tr>{th}</tr>{"".join(trs)}</table>'


def result_html(rows):
    head, body = rows[0], rows[1:]
    plan = head == ['QUERY PLAN']
    if len(body) > 12 and not plan:
        # long results wrap into two columns, so they stay legible on a slide
        half = math.ceil(len(body) / 2)
        tables = (f'<div style="display:inline-flex;gap:24px;align-items:flex-start">'
                  f'{result_table(head, body[:half], plan)}{result_table(head, body[half:], plan)}</div>')
    else:
        tables = result_table(head, body, plan)
    return page(f'<div id="card"><div class="label">Result</div>{tables}</div>')


class Renderer:
    """Collects HTML snippets, renders them all in one browser session."""

    def __init__(self):
        IMG_DIR.mkdir(parents=True, exist_ok=True)
        self.jobs = []

    def add(self, html):
        path = IMG_DIR / (hashlib.sha1(html.encode()).hexdigest()[:16] + '.png')
        if not path.exists() and all(j['file'] != str(path) for j in self.jobs):
            self.jobs.append({'file': str(path), 'html': html})
        return path

    def run(self):
        if not self.jobs:
            return
        jobs_file = IMG_DIR / 'jobs.json'
        jobs_file.write_text(json.dumps(self.jobs))
        subprocess.run(['node', str(ROOT / 'scripts/slides/render_examples.mjs'), str(jobs_file)],
                       check=True, stdout=subprocess.DEVNULL)
        jobs_file.unlink()
        self.jobs = []


def image_inches(path):
    """CSS px at 96 dpi, captured at 2x."""
    w, h = Image.open(path).size
    return w / 192, h / 192


# ---------------------------------------------------------------------------
# Text measurement (estimates — conservative, so boxes never overlap)
# ---------------------------------------------------------------------------

def text_lines(text, width_in, pt, mono=False):
    em = 0.60 if mono else 0.49
    per_line = max(10, int(width_in * 72 / (pt * em)))
    lines = 0
    for para in text.split('\n'):
        words, cur = para.split(), 0
        n = 1
        for w in words:
            if cur and cur + 1 + len(w) > per_line:
                n += 1
                cur = len(w)
            else:
                cur += (1 if cur else 0) + len(w)
        lines += n
    return lines


def block_height(b, width_in, pt):
    line = pt * 1.2 / 72
    gap = pt * 0.45 / 72
    if b.kind == 'list':
        return sum(text_lines(strip_inline(t), width_in - 0.35 - ind * 0.35, pt) * line + gap
                   for ind, _, t in b.items)
    if b.kind == 'quote':
        return text_lines(b.plain, width_in - 0.6, pt) * line + gap + 0.25
    return text_lines(b.plain, width_in, pt) * line + gap


def blocks_height(blocks, width_in, pt):
    return sum(block_height(b, width_in, pt) for b in blocks) + 0.1


# ---------------------------------------------------------------------------
# Slide building
# ---------------------------------------------------------------------------

class Deck:
    def __init__(self):
        self.prs = Presentation(str(TEMPLATE))
        self.tpl = list(self.prs.slides)

    # -- template plumbing -------------------------------------------------

    def clone(self, idx, keep_text=False):
        """New slide carrying the template slide's artwork (and text if asked)."""
        src = self.tpl[idx]
        slide = self.prs.slides.add_slide(src.slide_layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        rid_map = {rid: slide.part.relate_to(rel.target_part, RT.IMAGE)
                   for rid, rel in src.part.rels.items() if rel.reltype == RT.IMAGE}
        tree = slide.shapes._spTree
        for shp in src.shapes:
            is_text = shp.shape_type == MSO_SHAPE_TYPE.TEXT_BOX or (
                shp.is_placeholder and shp.shape_type != MSO_SHAPE_TYPE.PLACEHOLDER) or (
                shp.is_placeholder and not shp._element.xpath('.//p:blipFill'))
            if is_text and not keep_text:
                continue
            el = copy.deepcopy(shp._element)
            for node in el.iter():
                for attr in (qn('r:embed'), qn('r:link'), qn('r:id')):
                    if node.get(attr) in rid_map:
                        node.set(attr, rid_map[node.get(attr)])
            tree.insert_element_before(el, 'p:extLst')
        return slide

    def logo(self, slide):
        """The template's small mark, top left — on slides cloned from a blank."""
        src = self.tpl[T_TEXT]
        pic = next(s for s in src.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE and s.width < Inches(1))
        rid = slide.part.relate_to(src.part.rels[pic._element.blipFill.blip.rEmbed].target_part, RT.IMAGE)
        el = copy.deepcopy(pic._element)
        el.blipFill.blip.rEmbed = rid
        slide.shapes._spTree.insert_element_before(el, 'p:extLst')

    def finish(self, path):
        ids = self.prs.slides._sldIdLst
        for sld in list(ids)[:len(self.tpl)]:
            self.prs.part.drop_rel(sld.rId)
            ids.remove(sld)
        self.prs.save(str(path))

    # -- text primitives ----------------------------------------------------

    @staticmethod
    def textbox(slide, x, y, w, h, name=None, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        if name:
            tb.name = name
        return tb

    @staticmethod
    def runs(p, text, pt, color=INK, bold=False, italic=False):
        for t, st in parse_inline(text, bold, italic):
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(pt if not st['code'] else pt * 0.95)
            f.bold = st['bold']
            f.italic = st['italic']
            f.name = MONO if st['code'] else FONT
            f.color.rgb = ACCENT_TEXT if st['code'] else color
            if st['link']:
                r.hyperlink.address = st['link']
                f.color.rgb = ACCENT_TEXT
                f.underline = True

    @staticmethod
    def bullet(p, indent, ordered=False, start=1, pt=BODY_PT, char='•'):
        pPr = p._p.get_or_add_pPr()
        mar = int(Inches(0.32 + indent * 0.35))
        pPr.set('marL', str(mar))
        pPr.set('indent', str(-int(Inches(0.32))))
        for tag in ('a:buFont', 'a:buChar', 'a:buAutoNum', 'a:buNone'):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        font = pPr.makeelement(qn('a:buFont'), {'typeface': FONT})
        pPr.append(font)
        if ordered:
            pPr.append(pPr.makeelement(qn('a:buAutoNum'), {'type': 'arabicPeriod', 'startAt': str(start)}))
        else:
            pPr.append(pPr.makeelement(qn('a:buChar'), {'char': char}))

    def write_blocks(self, tf, blocks, pt=BODY_PT):
        first = True
        for b in blocks:
            if b.kind == 'list':
                counters = {}
                for ind, ordered, text in b.items:
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    counters[ind] = counters.get(ind, b.start if ind == 0 else 0) + 1
                    self.bullet(p, ind, ordered, counters[ind], pt, '•' if ind == 0 else '–')
                    p.space_after = Pt(pt * 0.45)
                    self.runs(p, text, pt)
                continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(pt * 0.45)
            if b.kind == 'quote':
                pPr = p._p.get_or_add_pPr()
                pPr.set('marL', str(int(Inches(0.3))))
                p.space_before = Pt(6)
                self.runs(p, b.text, pt - 1, ACCENT_TEXT, italic=False)
            else:
                self.runs(p, b.text, pt)

    def title(self, slide, text, kicker, width):
        """Kicker + title at the same spot on every content slide. Returns the y below it."""
        if kicker:
            tb = self.textbox(slide, LEFT, 0.28, width, 0.32, 'Kicker')
            p = tb.text_frame.paragraphs[0]
            self.runs(p, kicker.upper(), 12, ACCENT_TEXT, bold=True)
        pt = 36
        while text_lines(strip_inline(text), width, pt) > 1 and pt > 28:
            pt -= 4
        lines = text_lines(strip_inline(text), width, pt)
        h = lines * pt * 1.15 / 72 + 0.1
        tb = self.textbox(slide, LEFT, 0.62, width, h, 'Title')
        self.runs(tb.text_frame.paragraphs[0], text, pt, bold=True)
        return 0.62 + h + 0.2

    # -- slide types ----------------------------------------------------------

    def title_slide(self, title, subtitle):
        s = self.clone(T_TITLE)
        m = re.match(r'(Lecture \d+)\.\s*(.*)', title)
        lead, name = (m.group(1), m.group(2)) if m else ('', title)
        tb = self.textbox(s, 0.38, 1.25, 6.2, 4.6, 'Title')
        tf = tb.text_frame
        p = tf.paragraphs[0]
        self.runs(p, f'{COURSE}  ·  {lead}'.upper(), 16, ACCENT_TEXT, bold=True)
        p.space_after = Pt(18)
        p = tf.add_paragraph()
        pt = 54 if len(name) < 24 else 44
        self.runs(p, name, pt, bold=True)
        p.line_spacing = 0.95
        p.space_after = Pt(28)
        p = tf.add_paragraph()
        self.runs(p, subtitle, 20)

    def section_slide(self, title, intro, items, variant):
        if variant == 0:     # text left, artwork right
            s = self.clone(T_SECTION_A)
            x, w, y = 0.6, 5.6, 1.3
        else:                # artwork left, text right
            s = self.clone(T_SECTION_B)
            x, w, y = 6.6, 5.9, 0.75
        tb = self.textbox(s, x, y, w, 1.6, 'Title')
        pt = 40 if text_lines(title, w, 40) <= 2 else 32
        self.runs(tb.text_frame.paragraphs[0], title, pt, bold=True)
        y += text_lines(title, w, pt) * pt * 1.15 / 72 + 0.4
        blocks = list(intro)
        if items:
            blocks.append(Block('list', items=[[0, False, t] for t in items]))
        pt = 18
        while blocks and blocks_height(blocks, w, pt) > BOTTOM - y and pt > 14:
            pt -= 1
        if blocks:
            tb = self.textbox(s, x, y, w, BOTTOM - y, 'Body')
            self.write_blocks(tb.text_frame, blocks, pt)
        return s

    def text_slides(self, topic_title, kicker, blocks):
        """One or more slides of prose. Short ones get the template's artwork."""
        narrow_w = 7.4
        wide_w = RIGHT - LEFT
        if blocks_height(blocks, narrow_w, BODY_PT) <= BOTTOM - 2.25 and \
                text_lines(strip_inline(topic_title), 5.0, 36) <= 1 and \
                text_lines(kicker.upper(), 5.6, 12) <= 1:
            s = self.clone(T_TEXT)
            y = self.title(s, topic_title, kicker, 6.6)
            tb = self.textbox(s, LEFT, max(y, 2.25), narrow_w, BOTTOM - max(y, 2.25), 'Body')
            self.write_blocks(tb.text_frame, blocks)
            return
        # paginate on the plain slide
        pages, cur = [], []
        avail = BOTTOM - 1.75
        for b in self.split_long(blocks, wide_w, avail):
            if cur and blocks_height(cur + [b], wide_w, BODY_PT) > avail:
                pages.append(cur)
                cur = []
            cur.append(b)
        if cur:
            pages.append(cur)
        for a, b in zip(pages, pages[1:]):
            if len(a) > 1 and introduces(a[-1]):
                b.insert(0, a.pop())
        for blocks_on_page in pages:
            s = self.clone(T_BLANK)
            self.logo(s)
            y = self.title(s, topic_title, kicker, wide_w)
            tb = self.textbox(s, LEFT, y, wide_w, BOTTOM - y, 'Body')
            self.write_blocks(tb.text_frame, blocks_on_page)

    @staticmethod
    def split_long(blocks, width, avail):
        """Break lists that would not fit one slide into smaller lists."""
        out = []
        for b in blocks:
            if b.kind == 'list' and block_height(b, width, BODY_PT) > avail:
                chunk, start = [], 0
                for i, item in enumerate(b.items):
                    trial = Block('list', items=chunk + [item])
                    if chunk and block_height(trial, width, BODY_PT) > avail:
                        nb = Block('list', items=chunk)
                        nb.start = start
                        out.append(nb)
                        start, chunk = i, []
                    chunk.append(item)
                nb = Block('list', items=chunk)
                nb.start = start
                out.append(nb)
            else:
                out.append(b)
        return out

    def table_slides(self, topic_title, kicker, rows, lead=(), note=()):
        head, body = rows[0], rows[1:]
        width = RIGHT - LEFT
        ncol = len(head)
        # column widths from content length, with a floor
        lens = [max(len(strip_inline(r[c])) if c < len(r) else 0 for r in rows) for c in range(ncol)]
        # never narrower than the longest word, so nothing breaks mid-word
        def word_in(cell):
            # inches the longest word needs; inline code is set in the wider mono face
            return max([len(w) * TABLE_PT * (0.62 if t[1]['code'] else 0.55) / 72
                        for t in parse_inline(cell) for w in t[0].split()] or [0])
        floor = [max(word_in(r[c]) for r in rows if c < len(r)) + 0.25 for c in range(ncol)]
        weights = [min(max(l, 6), 60) for l in lens]
        tot = sum(weights)
        widths = [width * w_ / tot for w_ in weights]
        # lift narrow columns to their floor, taking the space from the others
        for _ in range(5):
            short = [c for c in range(ncol) if widths[c] < floor[c]]
            if not short:
                break
            need = sum(floor[c] - widths[c] for c in short)
            rest = [c for c in range(ncol) if c not in short]
            rest_w = sum(widths[c] for c in rest)
            for c in short:
                widths[c] = floor[c]
            for c in rest:
                widths[c] -= need * widths[c] / rest_w
        pad = 0.2

        def row_h(r, pt=TABLE_PT, bold=False):
            n = max(text_lines(strip_inline(r[c] if c < len(r) else ''), widths[c] - pad, pt) for c in range(ncol))
            return n * pt * 1.2 / 72 + 0.1

        pages = []
        while body or not pages:
            s = self.clone(T_BLANK)
            self.logo(s)
            y = self.title(s, topic_title, kicker, width)
            if lead and not pages:
                h = blocks_height(lead, width, CAPTION_PT)
                tb = self.textbox(s, LEFT, y, width, h, 'Lead')
                self.write_blocks(tb.text_frame, lead, CAPTION_PT)
                y += h + 0.1
            avail = BOTTOM - y - row_h(head)
            take, used = 0, 0
            while take < len(body) and used + row_h(body[take]) <= avail:
                used += row_h(body[take])
                take += 1
            take = max(take, 1) if body else 0
            rest = body[take:]
            if 0 < len(rest) <= 2 and used + sum(row_h(r) for r in rest) <= avail + 0.6:
                take = len(body)
            # avoid a lonely last row on the next slide
            if 0 < len(body) - take < 2 and take > 3:
                take -= 1
            chunk, body = body[:take], body[take:]
            pages.append(chunk)
            bottom = self.native_table(s, [head] + chunk, widths, y)
            if not body:
                break
        if note:
            h = blocks_height(note, width, CAPTION_PT)
            if bottom + 0.4 + h <= BOTTOM:
                tb = self.textbox(s, LEFT, bottom + 0.4, width, h, 'Note')
                self.write_blocks(tb.text_frame, note, CAPTION_PT)
            else:
                self.text_slides(topic_title, kicker, list(note))

    def native_table(self, slide, rows, widths, y):
        heights = []
        for r in rows:
            n = max(text_lines(strip_inline(r[c] if c < len(r) else ''), widths[c] - 0.2, TABLE_PT)
                    for c in range(len(widths)))
            heights.append(n * TABLE_PT * 1.2 / 72 + 0.1)
        shape = slide.shapes.add_table(len(rows), len(widths), Inches(LEFT), Inches(y),
                                       Inches(sum(widths)), Inches(sum(heights)))
        shape.name = 'Table'
        tbl = shape.table
        # plain style: no banding from the theme, we colour it ourselves
        tblPr = tbl._tbl.tblPr
        tblPr.set('bandRow', '0')
        tblPr.set('firstRow', '1')
        for c, w_ in enumerate(widths):
            tbl.columns[c].width = Inches(w_)
        for r, row in enumerate(rows):
            tbl.rows[r].height = Inches(heights[r])
            for c in range(len(widths)):
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.1)
                cell.margin_top = cell.margin_bottom = Inches(0.05)
                cell.fill.solid()
                cell.fill.fore_color.rgb = ACCENT if r == 0 else (ZEBRA if r % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                self.runs(p, row[c] if c < len(row) else '', TABLE_PT, INK, bold=(r == 0))
        return y + sum(heights)

    def example_slides(self, topic_title, kicker, lead, code_imgs, result_img, note, code_text):
        """Code (and result) images, laid out to stay as large as possible.

        code_imgs holds alternative renderings of the same code (one pane, two
        panes); the one that ends up largest on the slide wins.
        """
        if isinstance(code_imgs, list) and len(code_imgs) > 1:
            best = max(code_imgs, key=lambda c: self.example_scale(lead, c, result_img))
            return self.example_slides(topic_title, kicker, lead, best, result_img, note, code_text)
        code_img = code_imgs[0] if isinstance(code_imgs, list) else code_imgs
        width = RIGHT - LEFT
        imgs = [image_inches(code_img)] + ([image_inches(result_img)] if result_img else [])
        gap = 0.25

        def layouts(top):
            avail_h = BOTTOM - top
            opts = []
            # stacked: code above result
            sw = max(i[0] for i in imgs)
            sh = sum(i[1] for i in imgs) + gap * (len(imgs) - 1)
            opts.append(('stack', min(1.0, width / sw, avail_h / sh)))
            if len(imgs) == 2:
                sw = imgs[0][0] + imgs[1][0] + gap
                sh = max(imgs[0][1], imgs[1][1])
                opts.append(('side', min(1.0, width / sw, avail_h / sh)))
            return opts

        lead_blocks = list(lead)
        top0 = 0.62 + 0.55 + 0.2   # one-line title
        lead_h = blocks_height(lead_blocks, width, CAPTION_PT) if lead else 0
        top = top0 + lead_h + (0.1 if lead else 0)
        opts = layouts(top)
        # lead in a left column, images on the right — buys height for tall code
        col_w = 3.4
        if lead:
            right_w = width - col_w - 0.4
            sw = max(i[0] for i in imgs)
            sh = sum(i[1] for i in imgs) + gap * (len(imgs) - 1)
            opts.append(('column', min(1.0, right_w / sw, (BOTTOM - top0) / sh)))
        best = max(s for _, s in opts)
        # prefer the simplest layout that is nearly as large as the best
        kind = next(k for k, s_ in opts if s_ >= best * 0.92)
        scale = dict(opts)[kind]

        if result_img and best < 0.75:
            # too much for one slide: code, then result
            self.example_slides(topic_title, kicker, lead, [code_img], None, (), code_text)
            self.example_slides(topic_title, kicker, (), [result_img], None, note, code_text)
            return

        s = self.clone(T_BLANK)
        self.logo(s)
        y = self.title(s, topic_title, kicker, width)
        if lead:
            if kind == 'column':
                h = blocks_height(lead_blocks, col_w, CAPTION_PT)
                tb = self.textbox(s, LEFT, y, col_w, h, 'Lead')
            else:
                tb = self.textbox(s, LEFT, y, width, lead_h, 'Lead')
                y += lead_h + 0.1
            self.write_blocks(tb.text_frame, lead_blocks, CAPTION_PT)

        paths = [code_img] + ([result_img] if result_img else [])
        x0 = LEFT + col_w + 0.4 if kind == 'column' else LEFT
        if kind == 'column':
            y = max(y, top0) if not lead else y
        bottom_used = y
        if kind == 'side':
            x = x0
            for pth, (w_, h_) in zip(paths, imgs):
                pic = s.shapes.add_picture(str(pth), Inches(x), Inches(y), Inches(w_ * scale), Inches(h_ * scale))
                pic._element.nvPicPr.cNvPr.set('descr', code_text)
                x += w_ * scale + gap
                bottom_used = max(bottom_used, y + h_ * scale)
        else:
            yy = y
            for pth, (w_, h_) in zip(paths, imgs):
                pic = s.shapes.add_picture(str(pth), Inches(x0), Inches(yy), Inches(w_ * scale), Inches(h_ * scale))
                pic._element.nvPicPr.cNvPr.set('descr', code_text)
                yy += h_ * scale + gap
            bottom_used = yy - gap
        s.notes_slide.notes_text_frame.text = code_text

        if note:
            note_w = width if kind != 'column' else width - col_w - 0.4
            h = blocks_height(note, note_w, CAPTION_PT)
            if bottom_used + 0.25 + h <= BOTTOM:
                tb = self.textbox(s, x0, bottom_used + 0.25, note_w, h, 'Note')
                self.write_blocks(tb.text_frame, note, CAPTION_PT)
            else:
                self.text_slides(topic_title, kicker, list(note))

    def example_scale(self, lead, code_img, result_img):
        width = RIGHT - LEFT
        imgs = [image_inches(code_img)] + ([image_inches(result_img)] if result_img else [])
        top = 0.62 + 0.55 + 0.2 + (blocks_height(list(lead), width, CAPTION_PT) + 0.1 if lead else 0)
        sw = max(i[0] for i in imgs)
        sh = sum(i[1] for i in imgs) + 0.25 * (len(imgs) - 1)
        opts = [min(1.0, width / sw, (BOTTOM - top) / sh)]
        if len(imgs) == 2:
            opts.append(min(1.0, width / (imgs[0][0] + imgs[1][0] + 0.25), (BOTTOM - top) / max(i[1] for i in imgs)))
        if lead:
            opts.append(min(1.0, (width - 3.8) / sw, (BOTTOM - 1.37) / sh))
        return max(opts)

    def closing(self):
        self.clone(T_QA, keep_text=True)
        self.clone(T_THANKS, keep_text=True)


# ---------------------------------------------------------------------------
# Topic → slides
# ---------------------------------------------------------------------------

LEAD_MAX = 450      # characters of prose that may introduce an example
NOTE_MAX = 400      # characters of prose that may comment on one


def chars(b):
    return len(b.plain) if b.kind != 'list' else sum(len(strip_inline(t)) for _, _, t in b.items)


def introduces(b):
    return b.kind != 'list' and b.plain.rstrip().endswith(':')


def split_run(run, has_prev, has_next):
    """Split the prose between two examples into (note, middle, lead).

    The note comments on the previous example and goes on its slide; the lead
    introduces the next one. Whatever is left becomes slides of its own.
    """
    lead, note = [], []
    if has_next:
        total = 0
        while run and run[-1].kind in ('para', 'quote', 'list') and total + chars(run[-1]) <= LEAD_MAX:
            # a plain remark right after an example belongs to it, not to the next
            if has_prev and len(run) == 1 and not introduces(run[0]) and not lead:
                break
            total += chars(run[-1])
            lead.insert(0, run.pop())
            if not introduces(lead[0]) and lead[0].kind == 'para':
                break
    if has_prev:
        total = 0
        while run and not introduces(run[0]) and total + chars(run[0]) <= NOTE_MAX:
            total += chars(run[0])
            note.append(run.pop(0))
    return note, run, lead


def topic_slides(plan, renderer, title, kicker, blocks):
    """Append the slides for one topic to the plan."""
    # alternate prose runs and anchors (examples, tables)
    runs, anchors, cur = [], [], []
    for b in blocks:
        if b.kind in ('code', 'table'):
            runs.append(cur)
            anchors.append(b)
            cur = []
        else:
            cur.append(b)
    runs.append(cur)

    leads = [[] for _ in anchors]
    notes = [[] for _ in anchors]
    middles = []
    for k, run in enumerate(runs):
        note, middle, lead = split_run(list(run), k > 0, k < len(anchors))
        if k > 0:
            notes[k - 1] = note
        if k < len(anchors):
            leads[k] = lead
        middles.append(middle)

    for k, middle in enumerate(middles):
        if middle:
            plan.append(('text_slides', title, kicker, middle))
        if k == len(anchors):
            break
        b = anchors[k]
        if b.kind == 'table':
            plan.append(('table_slides', title, kicker, b.rows, leads[k], notes[k]))
        else:
            code_imgs = [renderer.add(code_html(b.text, b.lang))]
            split = code_html_split(b.text, b.lang)
            if split:
                code_imgs.append(renderer.add(split))
            result_img = renderer.add(result_html(b.result)) if b.result else None
            plan.append(('example_slides', title, kicker, leads[k], code_imgs, result_img, notes[k], b.text))


def build(md_path, renderer):
    blocks = parse_blocks(md_path.read_text())
    title, contents, sections = outline(blocks)

    # The plan is a list of (Deck method, *args). Building it first lets every
    # image be rendered in one browser session before any slide is laid out.
    plan = [('title_slide', title, AUTHOR),
            ('section_slide', 'Contents', [], [s.title for s in sections], 1)]
    agenda_intro = [b for b in contents if b.kind != 'list']
    if agenda_intro:
        plan.append(('text_slides', 'Contents', title, agenda_intro))

    for n, sec in enumerate(sections):
        intro = list(sec.intro)
        if not sec.topics:
            topic_slides(plan, renderer, sec.title, title, intro)
            continue
        # a short opening paragraph fits on the divider itself
        divider_intro = []
        if intro and intro[0].kind == 'para' and len(intro[0].plain) <= 260:
            divider_intro = [intro.pop(0)]
        plan.append(('section_slide', sec.title, divider_intro, [t.title for t in sec.topics], n % 2))
        topic_slides(plan, renderer, sec.title, title, intro)
        for t in sec.topics:
            topic_slides(plan, renderer, t.title, t.kicker, t.blocks)
    plan.append(('closing',))

    renderer.run()

    deck = Deck()
    for name, *args in plan:
        getattr(deck, name)(*args)
    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / md_path.with_suffix('.pptx').name
    deck.finish(out)
    return out


def export_pdf(pptx):
    pdf = pptx.with_suffix('.pdf')
    subprocess.run(['bash', str(ROOT / 'scripts/slides/export_pdf.sh'), str(pptx), str(pdf)], check=True)
    return pdf


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('lectures', nargs='*', default=LECTURES)
    ap.add_argument('--pdf', action='store_true', help='also export PDFs through PowerPoint')
    args = ap.parse_args()
    renderer = Renderer()
    for name in args.lectures:
        out = build(ROOT / name, renderer)
        print(f'wrote {out.relative_to(ROOT)}  ({len(Presentation(str(out)).slides)} slides)')
        if args.pdf:
            print(f'wrote {export_pdf(out).relative_to(ROOT)}')


if __name__ == '__main__':
    sys.exit(main())
